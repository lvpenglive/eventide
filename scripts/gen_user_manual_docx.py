# -*- coding: utf-8 -*-
"""Generate Eventide detailed user manual (Word)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm, Inches

SHOTS = Path("docs/screenshots")


def set_run_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = color


def add_figure(doc, filename, caption, width_in=6.2):
    """Embed screenshot if present; otherwise print a placeholder note."""
    path = SHOTS / filename
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.exists():
        doc.add_picture(str(path), width=Inches(width_in))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cap.add_run(caption)
        set_run_font(run, size=9, color=RGBColor(0x66, 0x66, 0x66))
        cap.paragraph_format.space_after = Pt(12)
    else:
        run = cap.add_run(f"【截图待补】{caption}（文件：screenshots/{filename}）")
        set_run_font(run, size=9, color=RGBColor(0x99, 0x55, 0x00))
        cap.paragraph_format.space_after = Pt(8)


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return p


def p(doc, text, size=11, bold=False, after=6):
    para = doc.add_paragraph()
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.35
    return para


def bullets(doc, items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(it)
        set_run_font(run, size=11)
        para.paragraph_format.space_after = Pt(2)


def numbered(doc, items):
    for it in items:
        para = doc.add_paragraph(style="List Number")
        run = para.add_run(it)
        set_run_font(run, size=11)
        para.paragraph_format.space_after = Pt(2)


def code(doc, text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)
    para.paragraph_format.left_indent = Cm(0.4)
    para.paragraph_format.space_after = Pt(8)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = header
        for para in cell.paragraphs:
            for run in para.runs:
                set_run_font(run, size=10, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=9)
    doc.add_paragraph()


def tip(doc, text):
    para = doc.add_paragraph()
    run = para.add_run("提示：")
    set_run_font(run, size=11, bold=True, color=RGBColor(0x1A, 0x56, 0xDB))
    run2 = para.add_run(text)
    set_run_font(run2, size=11)
    para.paragraph_format.space_after = Pt(8)


def warn(doc, text):
    para = doc.add_paragraph()
    run = para.add_run("注意：")
    set_run_font(run, size=11, bold=True, color=RGBColor(0xB4, 0x53, 0x09))
    run2 = para.add_run(text)
    set_run_font(run2, size=11)
    para.paragraph_format.space_after = Pt(8)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Eventide")
    set_run_font(r, size=28, bold=True, color=RGBColor(0x1A, 0x36, 0x5D))

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("用户使用手册")
    set_run_font(r2, size=22, bold=True)

    t3 = doc.add_paragraph()
    t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = t3.add_run(
        "控制台操作全量说明（详细版）\n"
        "适用对象：值班、运维、平台管理员\n"
        "配套文档：实施部署文档"
    )
    set_run_font(r3, size=11, color=RGBColor(0x55, 0x55, 0x55))

    p(
        doc,
        "本手册按控制台菜单与常见业务路径编写，说明「做什么、点哪里、填什么、如何验收」。"
        "部署与中间件安装请参阅《Eventide 全量实施部署文档》。",
    )
    tip(
        doc,
        "插图均为当前环境真实控制台截图（登录后采集）。"
        "若界面改版，可运行 scripts/capture_console_screenshots.py 覆盖 docs/screenshots/ 后重新生成本手册。",
    )

    h(doc, "目录", 1)
    bullets(
        doc,
        [
            "1. 产品简介与阅读指引",
            "2. 登录、界面与权限",
            "3. 总览",
            "4. 告警运营：告警事件、静默策略",
            "5. 接入配置：数据源、告警规则、告警接入",
            "6. SNMP Trap、MIB 库、Trap 策略",
            "7. 通知渠道、通知日志、告警丰富",
            "8. 系统管理：用户、角色、部门、系统设置",
            "9. 工具：Kafka",
            "10. 典型业务闭环（分步骤）",
            "11. 字段说明与模板变量",
            "12. 常见问题 FAQ",
            "13. 快捷对照表",
        ],
    )

    # 1
    h(doc, "1. 产品简介与阅读指引", 1)
    h(doc, "1.1 Eventide 能做什么", 2)
    p(
        doc,
        "Eventide 是多数据源告警引擎：既可主动拉取 Prometheus / VictoriaMetrics / Loki 等指标做规则评估，"
        "也可接收 Alertmanager、自研系统、业务拨测、Kafka 总线、SNMP Trap 等外部已判定告警；"
        "统一完成去重、丰富、静默与多渠道通知。",
    )
    h(doc, "1.2 两条主路径", 2)
    table(
        doc,
        ["路径", "适用", "关键步骤"],
        [
            [
                "规则评估",
                "自有监控可 PromQL/LogQL 查询",
                "数据源 → 通知渠道 → 告警规则 → 试跑 → 告警事件",
            ],
            [
                "告警接入",
                "外部系统已产生告警",
                "通知渠道 → 告警接入 → 试推送/真实推送 → 告警事件",
            ],
        ],
    )
    tip(doc, "多数客户「先配通知渠道，再配接入或规则」，避免告警产生了却无处可发。")

    # 2
    h(doc, "2. 登录、界面与权限", 1)
    h(doc, "2.1 登录", 2)
    numbered(
        doc,
        [
            "浏览器打开控制台地址（如 http://主机:8080 或公司反代域名）。",
            "输入用户名、密码，点击「进入系统」。",
            "登录成功后进入有权限的首页（通常为总览）。",
        ],
    )
    warn(
        doc,
        "请勿把账号密码写在网址查询参数中（如 ?username=&password=）。系统已改为 POST 登录；"
        "若地址栏误带密码参数，刷新后会被清除。密码会出现在浏览器历史，存在泄露风险。",
    )
    add_figure(
        doc,
        "01-login.png",
        "图 2-1 登录页（真实截图）",
    )
    h(doc, "2.2 界面结构", 2)
    bullets(
        doc,
        [
            "左侧导航：功能菜单，可折叠分组（告警运营 / 接入配置 / 通知丰富 / 系统管理 / 工具）。",
            "顶部：当前页标题、说明、页面操作按钮（新建、刷新、帮助等）。",
            "主区域：列表、卡片、表单与图表。",
            "右下角：操作提示 Toast；弹窗：新建/编辑/试推送等。",
            "侧栏底部：当前用户、主题切换（浅色/深色/跟随系统）、退出登录。",
        ],
    )
    h(doc, "2.3 主题与侧栏", 2)
    bullets(
        doc,
        [
            "主题：登录页或侧栏主题开关可切换浅色、深色或跟随系统。",
            "侧栏：点击折叠按钮可收起为仅图标，便于大屏值班。",
        ],
    )
    h(doc, "2.4 权限与只读许可", 2)
    p(
        doc,
        "管理员可在「用户管理 / 权限管理 / 部门管理」配置角色权限。无权限的菜单会隐藏。"
        "若产品许可试用结束且未导入有效授权，系统可能进入只读宽限：可登录查看，但新建/保存/删除等写操作会失败并提示许可只读（HTTP 402）。"
        "此时请到「系统设置 → 产品许可」导入授权文件。",
    )

    # 3
    h(doc, "3. 总览", 1)
    p(doc, "菜单：总览。展示系统运行快照，默认约每 30 秒自动刷新，也可点「刷新」。")
    bullets(
        doc,
        [
            "告警计数：firing / pending / resolved 等，可点击跳转到告警事件并带上状态筛选。",
            "资源计数：数据源、规则、渠道、接入等数量；若关键配置缺失，可能出现引导提示。",
            "最近告警：近期事件列表，便于值班快速扫一眼。",
            "通知跳过等统计：了解风暴节流/降级是否在起作用（视部署版本展示项而定）。",
        ],
    )
    add_figure(doc, "02-overview.png", "图 3-1 总览（真实截图）")

    # 4
    h(doc, "4. 告警运营", 1)
    h(doc, "4.1 告警事件", 2)
    p(doc, "菜单：告警运营 → 告警事件。这是日常值班主页面。")
    h(doc, "4.1.1 筛选与视图", 3)
    bullets(
        doc,
        [
            "状态：firing（触发中）、resolved（已恢复）、pending 等。",
            "级别：critical / warning / info 等。",
            "来源：规则评估或各类 ingress（如 ingress:alertmanager、ingress:snmptrap）。",
            "关键词 / IP：按名称、摘要、IP 等检索。",
            "存储：若启用 Elasticsearch，可在 MySQL / ES 之间切换检索。",
            "视图：卡片或表格；支持分页与自动刷新（视页面选项）。",
        ],
    )
    h(doc, "4.1.2 详情里看什么", 3)
    bullets(
        doc,
        [
            "基本信息：名称、状态、级别、指纹 fingerprint、起止时间、数值 value。",
            "标签 labels：实例、IP、来源、丰富后写入的台账字段等。",
            "注解 annotations：摘要、详情、Trap varbinds 等。",
            "通知日志：本告警相关通知发送结果，可点开查看正文。",
        ],
    )
    tip(doc, "同一条告警的 fire 与 recover 依赖相同 fingerprint（或拨测 messageId）。对不齐会导致「只触发不恢复」或重复通知。")
    add_figure(doc, "03-alerts.png", "图 4-1 告警事件列表（真实截图）")

    h(doc, "4.2 静默策略", 2)
    p(doc, "菜单：告警运营 → 静默策略。在时间窗内对匹配标签的告警抑制通知（事件仍可入库，视引擎行为以实际为准）。")
    numbered(
        doc,
        [
            "点击新建，填写名称、开始/结束时间。",
            "配置匹配标签（如 alertname、ip、severity）。标签匹配方式以表单说明为准。",
            "保存后，落在时间窗且标签匹配的告警将不再发通知（或按产品规则跳过）。",
            "到期后自动失效；也可手动删除/停用。",
        ],
    )
    warn(doc, "静默是运维止血手段，勿长期大范围静默掩盖真实故障。值班交接时请注明有效静默。")

    # 5
    h(doc, "5. 接入配置", 1)
    h(doc, "5.1 数据源", 2)
    p(doc, "菜单：接入配置 → 数据源。供「告警规则」拉取评估，不是 HTTP Webhook 接入。")
    table(
        doc,
        ["类型", "URL 示例", "规则 expr 写什么"],
        [
            ["prometheus / VictoriaMetrics", "http://prom:9090", "PromQL，如 up == 0"],
            ["log / Loki", "http://loki:3100", "返回数值向量的 LogQL"],
            ["kafka（管道字段告警）", "broker:9092", "JSON 数值字段路径，如 latency_ms"],
        ],
    )
    numbered(
        doc,
        [
            "新建 → 选择类型 → 填写名称、URL、是否启用。",
            "Kafka 数据源还需 topic、mode（field/depth/count）、field、label_fields 等。",
            "保存后到「告警规则」中引用该数据源。",
        ],
    )

    h(doc, "5.2 告警规则", 2)
    p(doc, "菜单：接入配置 → 告警规则。周期查询数据源，比较阈值，进入 pending → firing → resolved。")
    bullets(
        doc,
        [
            "名称、数据源、表达式 expr、比较符（> >= < <= == !=）、阈值。",
            "for_seconds：持续满足条件多久才转 firing（防抖）。",
            "interval_seconds：评估周期。",
            "severity、附加 labels、绑定的通知渠道。",
            "启用开关；「试跑」可立即评估一次看结果。",
        ],
    )
    tip(doc, "试跑成功不代表一定会通知：还需渠道可用、未静默、未命中风暴节流。")

    h(doc, "5.3 告警接入", 2)
    p(
        doc,
        "菜单：接入配置 → 告警接入。接收外部已判定告警。"
        "页面上方「使用帮助」可展开四种类型说明，并支持「一键创建」。",
    )
    h(doc, "5.3.1 类型说明", 3)
    table(
        doc,
        ["类型", "场景", "怎么接"],
        [
            [
                "Alertmanager",
                "Prometheus / VM 回调",
                "Webhook 指到 /api/ingress/{id}/alertmanager，带 Bearer Token",
            ],
            [
                "Generic / 拨测",
                "自研系统、Jeecg probe-alert",
                "POST /generic 或 /push；拨测需 eventType、messageId 等；IP 用 alertIp",
            ],
            [
                "Kafka",
                "告警总线、Zabbix、Trap 等",
                "填 Brokers+Topic+Group；无 map_* 自动识别，有 map_* 按映射解析",
            ],
            [
                "SNMP Trap 样例",
                "Trap→Kafka→本接入",
                "一键创建预填 Topic 的 Kafka 接入；字段已与 Trap 写出对齐",
            ],
        ],
    )
    add_figure(
        doc,
        "05-ingress-help-live.png",
        "图 5-1 告警接入「使用帮助」与一键创建（真实截图）",
    )
    add_figure(doc, "04-ingress.png", "图 5-2 告警接入列表（真实截图）")
    h(doc, "5.3.2 操作步骤", 3)
    numbered(
        doc,
        [
            "建议先建好通知渠道。",
            "新建接入或使用帮助里的「一键创建」，填写名称、Token（强烈建议必填）、绑定渠道。",
            "Kafka 类型填写 brokers、topic、start、partitions、group_id；按需打开字段映射。",
            "保存后，在卡片上复制接入地址；HTTP 类用「试推送」验证。",
            "到「告警事件」确认出现 firing/resolved，并检查通知日志。",
        ],
    )
    warn(
        doc,
        "未配置 Token 的 HTTP 接入，任何知道 URL 的人都能推告警。生产环境务必配置 Token，"
        "请求头使用 Authorization: Bearer <token> 或 X-Eventide-Token。",
    )
    h(doc, "5.3.3 字段映射（Generic / Kafka）", 3)
    p(
        doc,
        "当对方 JSON 不是内置 Alertmanager/Generic/拨测格式时，在编辑页开启「自定义字段映射」。"
        "填写点分路径（可截取），例如 map_name、map_ip、map_status、map_fingerprint 等。"
        "控制台帮助表有完整字段说明。路径截取支持 |before: |after: |split: 等语法。",
    )
    h(doc, "5.3.4 拨测字段对照", 3)
    table(
        doc,
        ["拨测字段", "Eventide"],
        [
            ["eventType fire/recover", "status firing/resolved"],
            ["messageId", "fingerprint（fire/recover 对齐）"],
            ["bizchainName 等", "labels.alertname"],
            ["retMessage", "annotations.summary"],
            ["retTimeMs", "value"],
            ["alertIp", "labels.ip / alertIp / instance"],
        ],
    )

    # 6
    h(doc, "6. SNMP Trap、MIB 库、Trap 策略", 1)
    p(doc, "三者配合：Trap 服务收 UDP → 策略匹配生成告警 JSON → Kafka → 告警接入消费。MIB 用于 OID 可读化。")

    h(doc, "6.1 SNMP Trap 页", 2)
    p(doc, "菜单：接入配置 → SNMP Trap。依赖部署侧 eventide-trap 与主服务 [trap] 反代。")
    bullets(
        doc,
        [
            "查看 Trap 服务健康、收包/解析/Kafka 计数。",
            "集群心跳：多机 Trap 实例列表（需 Redis 心跳）。",
            "试推送：模拟一条 Trap/告警，验证到 Kafka 与告警事件的闭环。",
            "可查看与 SNMP Trap Ingress 相关的消费组积压（视版本）。",
        ],
    )
    tip(doc, "若页面提示 Trap 不可达：检查 eventide-trap 是否启动、api_url/api_token 是否一致、防火墙是否放行。")
    add_figure(doc, "08-trap.png", "图 6-1 SNMP Trap 页（真实截图）")

    h(doc, "6.2 MIB 库", 2)
    p(doc, "菜单：接入配置 → MIB 库。上传厂商 MIB，浏览 OID 树。需配置 RustFS/MinIO（S3）。")
    numbered(
        doc,
        [
            "上传 MIB 文件（支持常见文本 MIB）。",
            "在模块列表中选择，展开 OID 树查看节点。",
            "节点详情可对目标设备做 SNMPv2c Get（需填目标与 community，注意安全）。",
            "Trap 实例会热加载 MIB，用于解析友好名称。",
        ],
    )

    h(doc, "6.3 Trap 策略", 2)
    p(doc, "菜单：接入配置 → Trap 策略。按 Trap OID / 条件决定是否告警、级别、文案模板等。")
    bullets(
        doc,
        [
            "支持筛选、启用开关、新建/编辑/删除。",
            "模板变量可点选插入（如对端 IP、OID、varbind）。",
            "支持 xlsx 导入导出；导入可选「合并」或「忽略重复」。",
            "保存后经 Redis 推送到各 Trap 实例热加载。",
        ],
    )
    numbered(
        doc,
        [
            "在 Eventide 配好策略与（可选）MIB。",
            "告警接入中创建/使用「SNMP Trap 样例」Kafka 接入，Topic 与 Trap 写出一致。",
            "绑定通知渠道并启用。",
            "Trap 页试推送 → 告警事件确认 labels.ip 等字段。",
            "需要主机台账时：告警丰富匹配键用 ip（Trap 写的是 labels.ip）。",
        ],
    )

    # 7
    h(doc, "7. 通知与丰富", 1)
    h(doc, "7.1 通知渠道", 2)
    p(doc, "菜单：通知丰富 → 通知渠道。")
    table(
        doc,
        ["类型", "主要配置"],
        [
            ["webhook", "URL、可选 secret"],
            ["钉钉 / 飞书 / 企微", "Webhook URL、加签 secret、是否 @all 等"],
            ["自定义 HTTP", "方法、URL、头、Bearer、JSON/文本模板"],
        ],
    )
    bullets(
        doc,
        [
            "可配置 firing / resolved 两套正文模板；留空用默认纯文本。",
            "模板支持变量插入（告警名、级别、IP、摘要、标签等）。",
            "保存后可用「测试」验证渠道可达（视页面按钮）。",
            "渠道可被规则或接入多选绑定。",
        ],
    )
    add_figure(doc, "06-channels.png", "图 7-1 通知渠道（真实截图）")

    h(doc, "7.2 通知日志", 2)
    p(doc, "菜单：通知丰富 → 通知日志。按渠道、成功/失败、关键词筛选，查看每次发送结果与正文，用于排障「为什么没收到」。")

    h(doc, "7.3 告警丰富", 2)
    p(doc, "菜单：通知丰富 → 告警丰富。含「台账数据」与「丰富规则」两个子页。")
    h(doc, "7.3.1 台账数据", 3)
    numbered(
        doc,
        [
            "新建台账（如 trap_hosts、cmdb_hosts）。",
            "设定匹配键（如 ip、instance）。",
            "导入或录入行数据（主机名、机房、联系人等列）。",
        ],
    )
    h(doc, "7.3.2 丰富规则", 3)
    numbered(
        doc,
        [
            "新建规则：匹配条件（如 source=ingress:snmptrap、某 label）。",
            "选择引用的台账与匹配方式。",
            "配置写入：是否写回 labels、描述模板如何拼接台账字段。",
            "使用「试跑预览」验证；控制台支持点选/拖拽变量芯片。",
            "保存后，后续告警在通知前（及事件展示）带上丰富结果。",
        ],
    )
    tip(doc, "Trap 场景匹配键请用 ip；不要误用 instance，除非你的标签里确实有 instance。")

    # 8
    h(doc, "8. 系统管理", 1)
    h(doc, "8.1 用户管理", 2)
    bullets(
        doc,
        [
            "新建用户：用户名、显示名、密码、部门、角色。",
            "可重置密码、禁用/删除（以界面为准）。",
            "首次空库管理员来自配置文件种子账号，上线后应立即改密。",
        ],
    )
    h(doc, "8.2 权限管理（角色）", 2)
    p(doc, "为角色勾选权限点（如 alerts:read、ingress:write、settings:write）。用户通过角色获得菜单与 API 能力。")
    h(doc, "8.3 部门管理", 2)
    p(doc, "维护组织树，用户归属部门，便于权限与职责划分。")
    h(doc, "8.4 系统设置", 2)
    bullets(
        doc,
        [
            "外观：主题相关。",
            "抗告警风暴：节流间隔、窗口次数、聚合窗口、接入削峰等；保存后当前集群热生效。",
            "告警历史 / Elasticsearch：可选同步与检索后端。",
            "Trap Token：与 Trap HTTP 共享的 api_token；可生成、查看掩码、轮换。",
            "产品许可：查看试用/授权状态，导出授权申请，导入授权文件。",
            "运行信息：版本与基础运行参数（部分敏感连接信息可能仅管理员可见，以实际版本为准）。",
        ],
    )
    add_figure(doc, "07-settings.png", "图 8-1 系统设置（真实截图）")
    warn(
        doc,
        "抗风暴「开太猛」可能导致通知过少；「关太松」可能打爆群机器人。"
        "建议先用默认，再按通知日志与业务容忍度微调。",
    )

    # 9
    h(doc, "9. 工具：Kafka", 1)
    p(doc, "菜单：工具 → Kafka。轻量管理界面，方便联调，不是替代专业 Kafka 管控台。")
    bullets(
        doc,
        [
            "填写 Brokers，刷新 Topic 列表。",
            "查看分区、消费组积压。",
            "浏览消息、试写一条 JSON（验证 Ingress 消费）。",
            "创建/删除 Topic（需具备相应权限；生产慎用删除）。",
        ],
    )
    tip(doc, "Brokers/Topic 会缓存在浏览器本地，便于下次打开；请与接入配置中的值保持一致。")

    # 10
    h(doc, "10. 典型业务闭环（分步骤）", 1)
    h(doc, "10.1 路径 A：Prometheus 规则告警", 2)
    numbered(
        doc,
        [
            "通知渠道：新建钉钉/飞书/Webhook 并测试。",
            "数据源：新建 prometheus，URL 指向 Prometheus/VM。",
            "告警规则：写 PromQL、比较符、阈值、for、绑定渠道，试跑。",
            "等待调度或再次试跑，在告警事件查看 firing。",
            "恢复条件后确认 resolved 与恢复通知。",
        ],
    )
    h(doc, "10.2 路径 B：Alertmanager Webhook", 2)
    numbered(
        doc,
        [
            "建通知渠道。",
            "告警接入 → Alertmanager，设 Token，绑定渠道。",
            "复制 Webhook URL 到 Alertmanager receiver。",
            "控制台「试推送」→ 告警事件 → 再放真实告警。",
        ],
    )
    h(doc, "10.3 路径 C：业务拨测", 2)
    numbered(
        doc,
        [
            "告警接入 → Generic，Token，绑定渠道。",
            "拨测平台 POST 到 /generic 或 /push，Body 含 eventType/messageId/bizchainName/retMessage/alertIp。",
            "fire 与 recover 使用同一 messageId。",
            "告警事件确认名称、IP、恢复闭环。",
        ],
    )
    h(doc, "10.4 路径 D：SNMP Trap", 2)
    numbered(
        doc,
        [
            "确认 Trap 服务与 Kafka、（可选）MIB/策略已就绪。",
            "告警接入 → 一键 SNMP Trap 样例，绑定渠道。",
            "SNMP Trap 页试推送。",
            "告警事件确认；需要时配置台账丰富（匹配 ip）。",
            "设备侧配置 Trap 目标为 Trap 的 UDP 地址/VIP。",
        ],
    )

    # 11
    h(doc, "11. 字段说明与模板变量", 1)
    h(doc, "11.1 告警核心字段", 2)
    table(
        doc,
        ["字段", "含义"],
        [
            ["fingerprint", "去重与恢复对齐的键"],
            ["status", "firing / resolved / pending"],
            ["severity", "级别"],
            ["labels", "多维标签，含 alertname、ip 等"],
            ["annotations", "摘要、描述等文本"],
            ["value", "当前数值（可选）"],
            ["startsAt / endsAt", "开始与结束时间"],
        ],
    )
    h(doc, "11.2 通知模板常用变量", 2)
    p(
        doc,
        "具体变量名以渠道编辑页芯片为准，常见包括告警名、状态、级别、IP、实例、摘要、指纹、来源、开始时间，"
        "以及丰富后的台账字段。自定义 HTTP 渠道还可使用 JSON 模板。",
    )

    # 12
    h(doc, "12. 常见问题 FAQ", 1)
    table(
        doc,
        ["问题", "处理建议"],
        [
            ["登录失败", "确认账号密码；是否被禁用；服务是否可达"],
            ["菜单看不到", "角色缺少权限；联系管理员赋权"],
            ["试推送成功但无通知", "查渠道、静默、风暴节流、通知日志失败原因"],
            ["只有触发没有恢复", "fingerprint/messageId 是否一致；对端是否发送 recover"],
            ["Kafka 接入无告警", "Topic/分区/Group；消息格式；字段映射；用工具→Kafka 试写"],
            ["Trap 页连不上", "trap 进程、api_token、反代 api_url、网络"],
            ["MIB 上传失败", "RustFS/S3 配置与 bucket；权限与时钟"],
            ["丰富不生效", "匹配条件、台账匹配键、是否 write_labels；试跑预览"],
            ["保存提示许可只读", "系统设置导入有效授权或联系厂商续期"],
            ["告警太多刷屏", "调高风暴节流/聚合；优化规则；临时静默"],
        ],
    )

    # 13
    h(doc, "13. 快捷对照表", 1)
    table(
        doc,
        ["我想…", "去哪里"],
        [
            ["看当前故障", "告警运营 → 告警事件"],
            ["临时不通知某类告警", "静默策略"],
            ["接 Prometheus 自己算", "数据源 + 告警规则"],
            ["接 Alertmanager", "告警接入 → Alertmanager"],
            ["接拨测/自研 JSON", "告警接入 → Generic"],
            ["接 Kafka 总线", "告警接入 → Kafka；工具→Kafka 联调"],
            ["接机房 Trap", "Trap 策略/MIB + Trap 页 + Kafka 接入样例"],
            ["改通知文案", "通知渠道 → 模板"],
            ["补主机名机房", "告警丰富 → 台账 + 规则"],
            ["改密码/加人", "用户管理"],
            ["防告警风暴", "系统设置 → 抗告警风暴"],
            ["导入许可证", "系统设置 → 产品许可"],
        ],
    )

    h(doc, "附录：推荐日常值班顺序", 1)
    numbered(
        doc,
        [
            "打开总览，看 firing 数量是否异常。",
            "进入告警事件，按级别处理 critical。",
            "对已知变更窗口加短时静默。",
            "抽查通知日志，确认渠道健康。",
            "交接时同步未恢复告警与有效静默。",
        ],
    )

    h(doc, "修订记录", 1)
    table(
        doc,
        ["日期", "说明"],
        [
            ["2026-08-07", "首版：控制台用户使用手册（详细版）"],
            ["2026-08-07", "增补界面截图/示意图与截图采集脚本说明"],
            ["2026-08-07", "服务可用后替换为全量真实控制台截图"],
        ],
    )

    out_en = Path("docs") / "Eventide-User-Manual.docx"
    out_zh = Path("docs") / (
        "Eventide" + "\u7528\u6237\u4f7f\u7528\u624b\u518c" + ".docx"
    )
    out_en.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_en)
    doc.save(out_zh)
    print("wrote", out_en)
    print("wrote", out_zh.name.encode("unicode_escape").decode("ascii"))


if __name__ == "__main__":
    build()
