# -*- coding: utf-8 -*-
"""Generate Eventide full deployment Word document."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm


def set_run_font(run, size=11, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    if color:
        run.font.color.rgb = color


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return p


def add_para(doc, text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.35
    return p


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(2)


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        cells = table.rows[r_i + 1].cells
        for c_i, val in enumerate(row):
            cells[c_i].text = str(val)
            for p in cells[c_i].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=9)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Eventide")
    set_run_font(r, size=28, bold=True, color=RGBColor(0x1A, 0x36, 0x5D))

    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = t2.add_run("全量实施部署文档")
    set_run_font(r2, size=22, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rm = meta.add_run(
        "多数据源告警引擎 · 评估 · 去重 · 通知\n"
        "适用版本：仓库 main / CI 发布包（Linux · 麒麟 · Windows）\n"
        "文档性质：实施部署全量说明（含 Trap / MIB / 许可）"
    )
    set_run_font(rm, size=11, color=RGBColor(0x55, 0x55, 0x55))

    add_para(doc, "说明：本文面向实施与运维人员，覆盖依赖准备、二进制部署、配置项、单机/多机、Trap+MIB、安全加固、验收与排障。敏感凭据请使用本环境自有值，勿将生产口令写入共享文档。")

    # TOC-like outline
    add_heading_cn(doc, "目录", 1)
    add_bullets(
        doc,
        [
            "1. 产品与架构概述",
            "2. 部署模式选择",
            "3. 环境与依赖要求",
            "4. 获取发布包",
            "5. 依赖组件部署（MySQL / Redis / Kafka / RustFS）",
            "6. Eventide 主服务部署",
            "7. SNMP Trap 服务部署",
            "8. 控制台初始化与最小闭环",
            "9. 多实例与高可用",
            "10. 产品许可（试用与离线授权）",
            "11. 安全加固清单",
            "12. 验收测试",
            "13. 运维与排障",
            "14. 附录：配置项速查",
        ],
    )

    # 1
    add_heading_cn(doc, "1. 产品与架构概述", 1)
    add_para(
        doc,
        "Eventide 是用 Rust 实现的轻量级多数据源告警引擎。核心能力包括：Prometheus/Loki 等规则评估、Alertmanager/Generic/Kafka/拨测等告警接入、去重、丰富、静默、通知，以及可选的 SNMP Trap 接收与 MIB/策略管理。",
    )
    add_heading_cn(doc, "1.1 核心组件", 2)
    add_table(
        doc,
        ["组件", "进程/产物", "职责"],
        [
            ["eventide", "主服务二进制", "HTTP API、JWT 控制台、规则调度、Kafka Ingress、通知队列、MIB/策略 CRUD"],
            ["eventide-trap", "Trap 二进制", "UDP SNMP Trap 接收 → 归一化 → Kafka；策略/MIB 热加载"],
            ["eventide-license", "厂商工具", "离线签发 / 校验商业授权文件"],
            ["static/", "静态控制台", "随主服务 ServeDir 直出，无独立 Node 构建"],
        ],
    )
    add_heading_cn(doc, "1.2 依赖中间件", 2)
    add_table(
        doc,
        ["依赖", "是否必需", "用途"],
        [
            ["MySQL 8+", "必需", "配置、告警事件、IAM、Trap 策略元数据等持久化"],
            ["Redis 7+", "必需", "抗风暴节流/聚合状态、集群选主、Trap 策略快照与 Pub/Sub、Trap 心跳"],
            ["Kafka", "按场景", "Trap 写出 Topic；Kafka 类告警接入消费"],
            ["RustFS / MinIO（S3）", "按场景", "MIB 文件正文对象存储；不用 MIB 库可不配"],
        ],
    )
    add_heading_cn(doc, "1.3 数据流（简图）", 2)
    add_code(
        doc,
        "外部告警/拨测/AM ──HTTP──► eventide ──► MySQL + 通知渠道\n"
        "Prometheus/Loki  ◄──调度──  eventide\n"
        "设备 Trap ──UDP──► eventide-trap ──Kafka──► eventide(Kafka Ingress) ──► 告警/通知\n"
        "控制台 CRUD MIB ──► MySQL 元数据 + RustFS 正文；策略 ──Redis──► Trap 热加载",
    )

    # 2
    add_heading_cn(doc, "2. 部署模式选择", 1)
    add_table(
        doc,
        ["模式", "包含内容", "适用"],
        [
            ["A. 最小告警闭环", "MySQL + Redis + eventide", "规则评估、AM/Generic 接入、通知；POC / 内网"],
            ["B. + Kafka 接入", "A + Kafka", "总线汇聚、Zabbix/拨测 JSON 等"],
            ["C. + SNMP Trap", "B + eventide-trap（+ 可选 RustFS）", "设备 Trap 进告警"],
            ["D. 生产多实例", "C + 多 eventide/trap + LB", "水平扩展与 HA"],
        ],
    )
    add_para(doc, "建议：先按模式 A 打通登录与通知，再按需叠加 Kafka / Trap / RustFS，最后做多机。")

    # 3
    add_heading_cn(doc, "3. 环境与依赖要求", 1)
    add_heading_cn(doc, "3.1 操作系统", 2)
    add_bullets(
        doc,
        [
            "Linux x86_64（glibc，Ubuntu 等）—— 使用 CI artifact：eventide-linux-x86_64",
            "银河麒麟 V10 SP3 x86_64 —— 使用 eventide-kylin-v10-sp3-x86_64",
            "Windows x86_64（MSVC 构建）—— 使用 eventide-windows-x86_64",
        ],
    )
    add_heading_cn(doc, "3.2 主机资源建议", 2)
    add_table(
        doc,
        ["规模", "CPU / 内存", "说明"],
        [
            ["POC / 小规模", "2C / 4GB+", "单机 MySQL+Redis 可同机或 Docker"],
            ["生产单机", "4C / 8GB+", "MySQL/Redis 建议独立主机"],
            ["多实例", "每节点 2C+ / 4GB+", "中间件独立集群"],
        ],
    )
    add_heading_cn(doc, "3.3 网络端口", 2)
    add_table(
        doc,
        ["端口", "组件", "说明"],
        [
            ["8080/TCP", "eventide", "API + 控制台（可改 listen）"],
            ["8081/TCP", "eventide-trap", "Trap HTTP（健康检查、试推送、反代）"],
            ["1162/UDP", "eventide-trap", "SNMP Trap 接收（可改；标准常为 162）"],
            ["3306/TCP", "MySQL", "仅内网"],
            ["6379/TCP", "Redis", "仅内网；务必认证与 ACL"],
            ["9092/TCP", "Kafka", "按集群规划"],
            ["9000/TCP", "RustFS/MinIO", "S3 API；仅内网"],
        ],
    )

    # 4
    add_heading_cn(doc, "4. 获取发布包", 1)
    add_para(doc, "推荐使用 GitHub Actions 构建产物，避免在生产机现场编译。")
    add_bullets(
        doc,
        [
            "仓库 Actions → Build 工作流：push / PR / workflow_dispatch 上传 Artifacts（保留约 30 天）",
            "打 tag（如 v0.1.0）可汇总进 GitHub Release",
            "每个包内含：eventide、eventide-trap、eventide-license、static/、*.toml.example、keys/license_public.pem",
        ],
    )
    add_para(doc, "解压示例（Linux）：", bold=True)
    add_code(
        doc,
        "tar -xzf eventide-linux-x86_64.tar.gz\n"
        "cd eventide-linux-x86_64\n"
        "cp eventide.toml.example eventide.toml\n"
        "cp eventide-trap.toml.example eventide-trap.toml\n"
        "# 编辑连接信息后启动",
    )
    add_para(doc, "源码编译（开发机）：", bold=True)
    add_code(
        doc,
        "cargo build -p eventide-server -p eventide-trap -p eventide-license --release\n"
        "# 产物：target/release/eventide、eventide-trap、eventide-license\n"
        "# 需自备 static/ 与配置文件，与二进制同目录或改 static_dir",
    )

    # 5
    add_heading_cn(doc, "5. 依赖组件部署", 1)
    add_heading_cn(doc, "5.1 MySQL", 2)
    add_para(doc, "要求：MySQL 8+，字符集 utf8mb4。库名建议 eventide。首次启动主服务会自动 migrate 建表。")
    add_para(doc, "本地 Docker Compose（仓库自带）：", bold=True)
    add_code(doc, "docker compose up -d   # 启动 MySQL 8 + Redis 7")
    add_para(doc, "连接串示例：mysql://eventide:eventide@127.0.0.1:3306/eventide")
    add_para(doc, "生产要求：独立实例、强口令、仅业务网可达、定期备份；勿使用 root 远程暴露。")

    add_heading_cn(doc, "5.2 Redis", 2)
    add_para(
        doc,
        "用于通知节流/聚合状态、调度选主、Trap 策略快照与变更广播、Trap 实例心跳、storm 配置多机热更新等。生产必须设置密码/ACL，禁止公网裸奔。",
    )
    add_para(doc, "连接串示例：redis://:password@127.0.0.1:6379/")

    add_heading_cn(doc, "5.3 Kafka（模式 B/C）", 2)
    add_bullets(
        doc,
        [
            "Trap 默认写出 Topic：eventide.snmptrap（可改）",
            "分区数需与 Trap kafka_partitions、Kafka Ingress options.partitions 一致",
            "Eventide 以 consumer group 消费；多实例自动 rebalance",
            "控制台「工具 → Kafka」可辅助查看 Topic / 积压 / 试写",
        ],
    )

    add_heading_cn(doc, "5.4 RustFS / MinIO（MIB 场景）", 2)
    add_para(
        doc,
        "仅当需要控制台「MIB 库」上传与 Trap 热加载 MIB 正文时必需。兼容 S3 API。创建 bucket（如 eventide-mibs），将 endpoint、access_key、secret_key、bucket、region 同时写入 eventide.toml [trap] 与 eventide-trap.toml。",
    )
    add_bullets(
        doc,
        [
            "不做 MIB：可不配 S3，主告警能力不受影响",
            "做 Trap 但不上传 MIB：策略仍可用 MySQL/Redis；OID 友好名等能力受限",
            "Server 负责写对象；Trap 负责读与本地缓存（mib_cache_dir）",
        ],
    )

    # 6
    add_heading_cn(doc, "6. Eventide 主服务部署", 1)
    add_heading_cn(doc, "6.1 目录布局建议", 2)
    add_code(
        doc,
        "/opt/eventide/\n"
        "  eventide                 # 或 Windows: eventide.exe\n"
        "  eventide.toml\n"
        "  static/                  # 控制台资源（与包内一致）\n"
        "  keys/license_public.pem\n"
        "  data/mib-cache/          # 可选\n"
        "  logs/                    # 建议由 systemd/nssm 重定向",
    )
    add_heading_cn(doc, "6.2 配置要点（eventide.toml）", 2)
    add_bullets(
        doc,
        [
            "listen：生产建议内网或反代后仅监听必要地址",
            "mysql_url / redis_url：指向上一节依赖",
            "static_dir：默认 static，与二进制相对路径",
            "[auth]：首次空库用 username/password 种子管理员；上线务必修改 password 与 jwt_secret",
            "[storm]：抗告警风暴默认；也可在控制台「系统设置」覆盖（多机经 Redis 热同步）",
            "[cluster]：多实例时开启选主，仅 leader 跑规则调度/聚合 flush",
            "[trap]：api_url 指向 Trap HTTP；api_token 与 Trap 一致；s3_* 按需",
        ],
    )
    add_heading_cn(doc, "6.3 启动", 2)
    add_code(
        doc,
        "# Linux\n"
        "./eventide eventide.toml\n\n"
        "# Windows\n"
        ".\\eventide.exe eventide.toml",
    )
    add_para(doc, "systemd 单元示例（Linux）：", bold=True)
    add_code(
        doc,
        "[Unit]\nDescription=Eventide Alert Engine\nAfter=network.target mysql.service redis.service\n\n"
        "[Service]\nType=simple\nWorkingDirectory=/opt/eventide\n"
        "ExecStart=/opt/eventide/eventide /opt/eventide/eventide.toml\n"
        "Restart=on-failure\nRestartSec=5\n\n"
        "[Install]\nWantedBy=multi-user.target",
    )
    add_heading_cn(doc, "6.4 反向代理（建议）", 2)
    add_para(
        doc,
        "生产将 https://eventide.example.com 反代到 8080；TLS 在边缘终止。多实例时对 /api 与静态页做 HTTP 负载均衡即可（会话为 JWT，无粘滞要求）。",
    )

    # 7
    add_heading_cn(doc, "7. SNMP Trap 服务部署", 1)
    add_para(doc, "模式 C 需要。与主服务共用同一 MySQL、Redis；使用 Kafka 写出；MIB 场景共用 RustFS。")
    add_heading_cn(doc, "7.1 配置要点（eventide-trap.toml）", 2)
    add_bullets(
        doc,
        [
            "listen_http / listen_udp：HTTP 与 UDP Trap 地址",
            "api_token：必须与 eventide.toml [trap] api_token 一致；生产禁止为空（为空则 HTTP fail-open）",
            "kafka_brokers / kafka_topic / kafka_partitions：与集群及 Ingress 对齐",
            "mysql_url / redis_url / s3_*：与主服务一致",
            "instance_id / ha_vip / heartbeat_secs：多机 HA 时配置",
            "[[snmpv3_users]]：需要收 v3 时配置 USM",
        ],
    )
    add_heading_cn(doc, "7.2 启动与主服务联动", 2)
    add_code(doc, "./eventide-trap eventide-trap.toml")
    add_bullets(
        doc,
        [
            "主服务 [trap] api_url 指向本机或 VIP/LB 的 Trap HTTP",
            "控制台「SNMP Trap」经 /trap-api 反代到 Trap（注入 Bearer token）",
            "告警接入：创建 Kafka Ingress，Topic = Trap 写出 Topic；可用「一键 SNMP Trap 样例」",
        ],
    )
    add_heading_cn(doc, "7.3 策略分发", 2)
    add_para(
        doc,
        "MySQL trap_policies 为真相源 → Eventide CRUD 后发布 Redis 快照并 PUBLISH → Trap 订阅热加载；失败回退 MySQL；policy_reload_secs 轮询兜底。",
    )

    # 8
    add_heading_cn(doc, "8. 控制台初始化与最小闭环", 1)
    add_para(doc, "浏览器打开：http://<主机>:8080 （或反代域名）。默认种子账号见 [auth]（示例常为 admin，密码务必修改）。")
    add_heading_cn(doc, "8.1 推荐初始化顺序", 2)
    add_bullets(
        doc,
        [
            "登录 → 修改管理员密码（用户管理）",
            "通知渠道：钉钉 / 飞书 / 企微 / webhook / 自定义 HTTP",
            "（可选）数据源 + 告警规则；或告警接入 Alertmanager / Generic / Kafka",
            "试推送 → 告警事件确认 → 再接真实平台",
            "（可选）告警丰富 / 台账；静默策略；抗告警风暴设置",
            "（Trap）MIB 库 → Trap 策略 → Kafka 接入 → 试推送 / 真实 Trap",
        ],
    )
    add_heading_cn(doc, "8.2 Alertmanager 接入", 2)
    add_code(
        doc,
        "Webhook URL: POST http(s)://<eventide>/api/ingress/{id}/alertmanager\n"
        "Header: Authorization: Bearer <token>\n"
        "或 X-Eventide-Token: <token>",
    )
    add_heading_cn(doc, "8.3 Generic / 拨测", 2)
    add_para(
        doc,
        "POST /api/ingress/{id}/generic 或 /push。拨测 JSON 建议含 eventType、messageId、bizchainName、retMessage；IP 使用 alertIp。",
    )

    # 9
    add_heading_cn(doc, "9. 多实例与高可用", 1)
    add_heading_cn(doc, "9.1 Eventide 多实例", 2)
    add_bullets(
        doc,
        [
            "共用同一 MySQL + Redis + 相同 jwt_secret / 配置语义",
            "[cluster] 选主：仅 leader 跑规则调度与聚合 flush",
            "Kafka Ingress：同 group_id 自动分摊分区",
            "HTTP / Webhook：所有实例均可服务；前置 LB 即可",
            "抗风暴配置：控制台保存后经 Redis 同步热加载（无需逐台重启）",
            "控制台：任一台或经统一入口访问均可（数据在库中）",
        ],
    )
    add_heading_cn(doc, "9.2 Trap 多机", 2)
    add_para(doc, "推荐：UDP 负载均衡，多台 Trap 并行收包写入同一 Kafka Topic。备选：Keepalived VIP（同时仅 MASTER 收包）。详见仓库 deploy/trap-ha/。")

    # 10
    add_heading_cn(doc, "10. 产品许可", 1)
    add_bullets(
        doc,
        [
            "首次无商业证：自动 30 天试用（可写）",
            "过期：只读宽限，写操作返回 402 / license_readonly；仍可导入授权",
            "客户：系统设置 → 产品许可 → 导出授权申请 → 厂商签发 → 导入授权文件",
            "厂商：eventide-license issue --request … -o xxx.eventide-lic.json（私钥勿入客户环境）",
        ],
    )

    # 11
    add_heading_cn(doc, "11. 安全加固清单", 1)
    add_table(
        doc,
        ["优先级", "项", "建议"],
        [
            ["高", "配置与凭据", "eventide.toml / trap 配置勿提交 Git；轮换已暴露口令；用环境隔离"],
            ["高", "MySQL / Redis", "内网、强认证、禁公网；Redis 存有策略与 token 快照"],
            ["高", "默认账号", "修改种子管理员密码与 jwt_secret"],
            ["高", "Ingress Token", "所有 HTTP 接入强制配置 Token"],
            ["高", "Trap api_token", "生产禁止为空；与主服务一致"],
            ["中", "监听地址", "0.0.0.0 仅内网或改反代后本机监听"],
            ["中", "TLS", "边缘 HTTPS；控制台勿在公网明文"],
            ["中", "日志", "避免日志打印完整 DSN（含密码）"],
            ["中", "RBAC", "按部门/角色最小权限；定期审查"],
            ["低", "CORS", "生产收敛允许源（当前开发态较宽松）"],
        ],
    )
    add_para(doc, "注意：登录表单已避免账号密码进入 URL 查询串；请始终使用正式登录页，勿把密码写在书签链接中。")

    # 12
    add_heading_cn(doc, "12. 验收测试", 1)
    add_table(
        doc,
        ["编号", "检查项", "期望"],
        [
            ["V1", "打开控制台并登录", "进入总览，无 URL 泄露密码"],
            ["V2", "GET /api/health（若暴露）或登录后总览", "服务正常"],
            ["V3", "创建通知渠道并试发", "渠道可达"],
            ["V4", "创建 Ingress + 试推送", "告警事件出现；有 Token 时无 Token 被拒"],
            ["V5", "（可选）规则试跑", "firing/resolved 符合预期"],
            ["V6", "（Trap）HTTP health + 试推送", "Kafka 有消息且 Ingress 生成告警"],
            ["V7", "（MIB）上传 MIB / 策略", "Trap 热加载后匹配生效"],
            ["V8", "（多机）杀一实例", "LB 仍可访问；选主/消费组恢复"],
            ["V9", "许可导入", "状态正确；过期只读行为符合预期"],
        ],
    )

    # 13
    add_heading_cn(doc, "13. 运维与排障", 1)
    add_table(
        doc,
        ["现象", "排查方向"],
        [
            ["无法登录", "MySQL 连通；[auth] 种子是否已建用户；JWT secret 是否被改导致旧 token 失效"],
            ["Ingress 401", "Token 头是否正确；路由是否启用"],
            ["Kafka 无消费", "brokers/topic/group；分区数；网络 ACL"],
            ["Trap 无告警", "UDP 是否到达；Kafka 写出；Ingress 是否订阅同 Topic；策略是否丢弃"],
            ["MIB 上传失败", "S3 endpoint/bucket/密钥；时钟；bucket 是否存在"],
            ["Redis pubsub 抖动", "网络/防火墙空闲断开；Trap 会自动重连并轮询兜底"],
            ["风暴不生效", "系统设置是否保存；多机是否均已热加载；Redis 是否可用"],
            ["402 只读", "许可/试用状态；导入有效授权"],
        ],
    )
    add_para(doc, "日志：建议用 systemd/journald 或 Windows 服务包装收集 stdout/stderr；勿将含密钥的配置提交到文档库。")

    # 14
    add_heading_cn(doc, "14. 附录：配置项速查", 1)
    add_heading_cn(doc, "14.1 eventide.toml 主要字段", 2)
    add_table(
        doc,
        ["字段", "说明"],
        [
            ["listen", "HTTP 监听"],
            ["mysql_url", "MySQL 连接串"],
            ["redis_url", "Redis 连接串"],
            ["static_dir", "控制台静态目录"],
            ["scheduler_tick_seconds", "规则调度周期"],
            ["[auth].*", "种子管理员、JWT"],
            ["[storm].*", "节流/聚合/削峰默认"],
            ["[cluster].*", "多实例选主"],
            ["[trap].*", "Trap 反代与 S3/MIB"],
            ["[elasticsearch].*", "可选告警历史 ES"],
        ],
    )
    add_heading_cn(doc, "14.2 eventide-trap.toml 主要字段", 2)
    add_table(
        doc,
        ["字段", "说明"],
        [
            ["listen_http / listen_udp", "HTTP / UDP"],
            ["api_token", "Trap HTTP 鉴权"],
            ["kafka_*", "写出 Kafka"],
            ["mysql_url / redis_url", "同主服务库"],
            ["s3_* / mib_cache_dir", "S3/MIB"],
            ["policy_reload_secs", "策略轮询兜底秒数"],
            ["instance_id / ha_vip / heartbeat_secs", "多机注册"],
            ["[[snmpv3_users]]", "SNMPv3 USM"],
        ],
    )
    add_heading_cn(doc, "14.3 相关仓库路径", 2)
    add_bullets(
        doc,
        [
            "示例配置：eventide.toml.example、eventide-trap.toml.example",
            "本地依赖：docker-compose.yml",
            "Trap HA 示例：deploy/trap-ha/",
            "构建工作流：.github/workflows/build.yml",
            "详细产品说明：README.md",
        ],
    )

    add_heading_cn(doc, "修订记录", 1)
    add_table(
        doc,
        ["日期", "说明"],
        [
            ["2026-08-07", "首版：全量实施部署文档（含 RustFS/Trap/多机/许可/安全）"],
        ],
    )

    out_en = Path("docs") / "Eventide-Full-Deployment-Guide.docx"
    out_zh = Path("docs") / (
        "Eventide" + "\u5168\u91cf\u5b9e\u65bd\u90e8\u7f72\u6587\u6863" + ".docx"
    )
    out_en.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_en)
    doc.save(out_zh)
    print(f"wrote {out_en.resolve()}")
    print(f"wrote {out_zh.name.encode('unicode_escape').decode('ascii')}")


if __name__ == "__main__":
    build()
